import io
from typing import List
from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy.orm import Session
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from app import models
from app.database import get_db

router = APIRouter(prefix="/export", tags=["export"])

HEADERS = [
    ("Date", "measured_at"),
    ("Weight (kg)", "weight"),
    ("BMI", "bmi"),
    ("Body Fat (%)", "body_fat"),
    ("Muscle Mass (kg)", "muscle_mass"),
    ("Skeletal Muscle (%)", "skeletal_muscle"),
    ("Visceral Fat", "visceral_fat"),
    ("Subcutaneous Fat (%)", "subcutaneous_fat"),
    ("Protein (%)", "protein"),
    ("Body Water (%)", "body_water"),
    ("BMR (kcal)", "bmr"),
    ("Body Age", "body_age"),
    ("Standard Weight (kg)", "standard_weight"),
    ("Fat Mass (kg)", "fat_mass"),
    ("Fat Loss (kg)", "fat_loss"),
    ("Muscle Frequency", "muscle_frequency"),
    ("Skeletal Mass (kg)", "skeletal_mass"),
]

HEADER_FILL = PatternFill(start_color="4F81BD", end_color="4F81BD", fill_type="solid")
HEADER_FONT = Font(bold=True, color="FFFFFF")


@router.get("/profiles")
def export_profiles(
    profile_ids: str = Query(..., description="Comma-separated profile IDs"),
    db: Session = Depends(get_db),
):
    try:
        ids = [int(pid.strip()) for pid in profile_ids.split(",") if pid.strip()]
    except ValueError:
        raise HTTPException(status_code=400, detail="Invalid profile_ids parameter")

    if not ids:
        raise HTTPException(status_code=400, detail="No profile IDs provided")

    profiles = (
        db.query(models.Profile).filter(models.Profile.id.in_(ids)).all()
    )
    if not profiles:
        raise HTTPException(status_code=404, detail="No profiles found")

    wb = Workbook()
    wb.remove(wb.active)  # remove default empty sheet

    for profile in profiles:
        sheet_name = profile.name[:31]  # Excel sheet names max 31 chars
        ws = wb.create_sheet(title=sheet_name)

        # Profile info header block
        ws.append([f"Profile: {profile.name}"])
        ws.append([f"Height: {profile.height} cm", f"Age: {profile.age}", f"Gender: {profile.gender}"])
        ws.append([])

        # Column headers
        ws.append([h for h, _ in HEADERS])
        header_row = ws.max_row
        for col_idx in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=header_row, column=col_idx)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        # Measurement rows (newest first)
        measurements = (
            db.query(models.Measurement)
            .filter(models.Measurement.profile_id == profile.id)
            .order_by(models.Measurement.measured_at.desc())
            .all()
        )

        for m in measurements:
            row = []
            for _, attr in HEADERS:
                value = getattr(m, attr, None)
                if attr == "measured_at" and value is not None:
                    value = value.strftime("%Y-%m-%d %H:%M:%S")
                row.append(value)
            ws.append(row)

        # Auto-fit column widths
        for col_idx, (header, _) in enumerate(HEADERS, start=1):
            col_letter = get_column_letter(col_idx)
            ws.column_dimensions[col_letter].width = max(len(header) + 2, 14)

    output = io.BytesIO()
    wb.save(output)
    output.seek(0)

    return StreamingResponse(
        output,
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": "attachment; filename=femmto_export.xlsx"},
    )
